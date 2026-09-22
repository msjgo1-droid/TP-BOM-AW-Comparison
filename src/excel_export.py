# -*- coding: utf-8 -*-
"""비교 결과 + 체크리스트를 엑셀(xlsx)로 내보내기."""
import io
from typing import List, Dict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.i18n import t


def build_excel(rows: List[Dict], lang: str) -> bytes:
    """
    rows: [{
      "group": str, "category": str, "old_file": str, "new_file": str,
      "page": int|str, "status": str, "text_status": str,
      "confirmed": bool, "note": str,
    }, ...]
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    headers = [
        t(lang, "excel_col_group"),
        t(lang, "excel_col_category"),
        t(lang, "excel_col_old_file"),
        t(lang, "excel_col_new_file"),
        t(lang, "excel_col_page"),
        t(lang, "excel_col_status"),
        t(lang, "excel_col_text_status"),
        t(lang, "excel_col_confirmed"),
        t(lang, "excel_col_note"),
    ]
    ws.append(headers)
    header_fill = PatternFill(start_color="1F2A44", end_color="1F2A44", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        confirmed_text = t(lang, "excel_yes") if row.get("confirmed") else t(lang, "excel_no")
        ws.append([
            row.get("group", ""),
            row.get("category", ""),
            row.get("old_file", ""),
            row.get("new_file", ""),
            row.get("page", ""),
            row.get("status", ""),
            row.get("text_status", ""),
            confirmed_text,
            row.get("note", ""),
        ])

    widths = [22, 10, 34, 34, 8, 20, 20, 12, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
